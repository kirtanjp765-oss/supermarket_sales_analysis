# takes all the analysis dataframes and exports them to an excel file
# used openpyxl for some basic header styling

import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "outputs", "reports")
REPORT_PATH = os.path.join(REPORT_DIR, "supermarket_analysis_report.xlsx")


def style_header(ws, row_num, num_cols):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True, size=11)

    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill      = fill
        cell.font      = font
        cell.alignment = Alignment(horizontal="center")


def auto_column_width(ws):
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 40)


def export_to_excel(summary_stats, monthly_revenue, city_revenue,
                    product_perf, customer_analysis, payment_breakdown,
                    rfm, peak_hours):

    os.makedirs(REPORT_DIR, exist_ok=True)

    sheets = {
        "Summary Stats":       summary_stats,
        "Monthly Revenue":     monthly_revenue,
        "City Revenue":        city_revenue,
        "Product Performance": product_perf,
        "Customer Analysis":   customer_analysis,
        "Payment Methods":     payment_breakdown,
        "RFM Analysis":        rfm,
        "Peak Hours":          peak_hours,
    }

    with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)

        wb = writer.book

        for sheet_name, df in sheets.items():
            ws = wb[sheet_name]

            ws.cell(row=1, column=1, value=f"Supermarket Sales - {sheet_name}")
            ws.cell(row=1, column=1).font = Font(bold=True, size=12, color="1F4E79")

            style_header(ws, row_num=2, num_cols=len(df.columns))
            auto_column_width(ws)

    report_path = os.path.abspath(REPORT_PATH)
    print(f"excel report saved: {report_path}")
    return report_path
